"""Regresiones del aislamiento de los Excel entre cargas simultáneas."""

from concurrent.futures import ThreadPoolExecutor
from contextlib import contextmanager
from pathlib import Path
from threading import Barrier
import unittest
from unittest.mock import MagicMock, patch

from flask import Flask
from openpyxl import load_workbook

from views.subir_pedidos import routes


class ArchivosPedidosTests(unittest.TestCase):
    def setUp(self):
        self.app = Flask(__name__)
        self.app.config.update(TESTING=True, SECRET_KEY="solo-pruebas")
        self.app.register_blueprint(routes.subir_pedidos_bp)

    @staticmethod
    def pedidos(empresa):
        return {
            "rutas_con_placa": [{"ruta": 1, "placa": "ABC123"}],
            "data_ped": [
                {"ruta": 1, "codigo_pro": empresa, "producto": empresa, "pedir": 7}
            ],
            "total_rutas": 1,
        }

    @staticmethod
    @contextmanager
    def navegador():
        yield MagicMock()

    def cargar(self, empresa):
        with self.app.test_client() as client:
            with client.session_transaction() as session:
                session["user_id"] = 1
                session["empresa"] = empresa
            return client.post(
                "/subir-pedidos/login-portal",
                json={"usuario": "prueba", "contrasena": "prueba"},
            )

    def test_cargas_simultaneas_con_la_misma_ruta_conservan_su_excel(self):
        primera_carga_lista = Barrier(2)
        ambas_cargas_listas = Barrier(2)
        capturados = {}

        def obtener_pedidos(empresa):
            # La segunda carga empieza a escribir cuando el primer Excel ya existe.
            if empresa == "empresa_b":
                primera_carga_lista.wait(timeout=15)
            return self.pedidos(empresa)

        def subir(ruta_placa, archivo, campo_placa, **kwargs):
            from flask import session

            empresa = session["empresa"]
            if empresa == "empresa_a":
                primera_carga_lista.wait(timeout=15)
            ambas_cargas_listas.wait(timeout=15)
            with open(archivo, "rb") as excel:
                workbook = load_workbook(excel, read_only=True)
                try:
                    codigo = workbook.active.cell(5, 1).value
                finally:
                    workbook.close()
            capturados[empresa] = (archivo, codigo)
            return True

        with (
            patch.object(routes, "log_pedidos_rutas", side_effect=obtener_pedidos),
            patch.object(routes, "iniciar_navegador", self.navegador),
            patch.object(routes, "login_portal_grupo_nutresa", return_value=True),
            patch.object(routes, "cargar_pedido_masivo_excel", side_effect=subir),
            ThreadPoolExecutor(max_workers=2) as executor,
        ):
            respuestas = list(executor.map(self.cargar, ("empresa_a", "empresa_b")))

        for respuesta in respuestas:
            self.assertEqual(respuesta.status_code, 200)
            self.assertTrue(respuesta.json["success"])
        self.assertEqual(capturados["empresa_a"][1], "empresa_a")
        self.assertEqual(capturados["empresa_b"][1], "empresa_b")
        self.assertNotEqual(capturados["empresa_a"][0], capturados["empresa_b"][0])
        for archivo, _ in capturados.values():
            self.assertFalse(Path(archivo).exists())

    def test_limpia_archivos_si_falla_la_subida(self):
        archivos = []

        def fallar(ruta_placa, archivo, campo_placa, **kwargs):
            self.assertTrue(Path(archivo).exists())
            archivos.append(archivo)
            raise RuntimeError("Fallo simulado del portal")

        with (
            patch.object(routes, "log_pedidos_rutas", side_effect=self.pedidos),
            patch.object(routes, "iniciar_navegador", self.navegador),
            patch.object(routes, "login_portal_grupo_nutresa", return_value=True),
            patch.object(routes, "cargar_pedido_masivo_excel", side_effect=fallar),
        ):
            respuesta = self.cargar("empresa_a")

        self.assertEqual(respuesta.status_code, 500)
        self.assertFalse(respuesta.json["success"])
        self.assertEqual(len(archivos), 1)
        self.assertFalse(Path(archivos[0]).exists())


if __name__ == "__main__":
    unittest.main()
